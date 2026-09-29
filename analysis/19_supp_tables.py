"""Supplementary Tables 2-4 as one Excel workbook plus CSV copies, and Supplementary Table 1 as CSV.
Table 1: headline results by subgroup (claims_by_group.csv from step 11).
Table 2: all 344 port-cyclone events (panel_pairs.csv + het_port_traits.csv).
Table 3: pooled changes by zone, measure, period and cyclone class (fig2_zone_ratios.csv, divergence_pooled.csv).
Table 4: arrival-management savings by approach length, wait reduction and fuel price (jit_scenarios.csv).
Writes tables/supplementary_tables_2-4.xlsx, tables/supp_table{1,2,3,4}.csv in HURR_OUT.
"""
import os
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

OUT = os.environ.get("HURR_OUT", "out")
TAB = os.path.join(OUT, "tables"); os.makedirs(TAB, exist_ok=True)
pd.read_csv(os.path.join(OUT, "claims_by_group.csv")).to_csv(os.path.join(TAB, "supp_table1.csv"), index=False)
CLS = {"Tropical storm <50 kt": "Tropical storm (<50 kt)", "Strong TS 50–63 kt": "Strong tropical storm (50–63 kt)",
       "Hurricane ≥64 kt": "Hurricane (≥64 kt)", "All events": "All cyclones", "All storms": "All cyclones"}
CORD = list(dict.fromkeys(CLS.values()))
PER = {"prep": "2 days before", "storm": "Storm window", "after": "10 days after", "event30": "Whole event (−2 to +30 days)"}
ZONE = {"total": "Study region (all zones)", "port": "Berths", "channel": "Channels", "anchorage": "Anchorages",
        "offshore_hold": "Offshore holding areas", "offshore": "Open water"}
MEAS = {"hours": "Ship-hours", "co2": "CO2", "nox": "NOx", "pm25": "PM2.5"}

# ---------------- Table 2
P = pd.read_csv(os.path.join(OUT, "panel_pairs.csv"))
PT = pd.read_csv(os.path.join(OUT, "het_port_traits.csv")).set_index("port")
t2 = P.assign(port_type=P.port.map(PT.port_type), coast=P.port.map(PT.region)).sort_values(["season", "SID", "port"]).reset_index(drop=True)
T2 = pd.DataFrame({
    "Event": range(1, len(t2) + 1), "Cyclone": t2.storm.str.title(), "IBTrACS ID": t2.SID, "Season": t2.season.astype(str), "Port": t2.port,
    "Port type": t2.port_type, "Coast": t2.coast, "Class": t2.intensity.map(CLS), "Wind at closest approach (kt)": t2.wind_at_closest_kt,
    "Saffir–Simpson category": t2.sshs, "Closest distance (km)": t2.min_dist_km.round(1), "Landfall": t2.landfall.map({1: "Yes", 0: "No"}),
    "Storm surge (m)": t2.surge_m.round(2), "Rainfall (mm)": t2.rain_mm.round(1), "Storm window (days)": t2.k_days, "Baseline years": t2.n_base})

# ---------------- Table 3
Z = pd.read_csv(os.path.join(OUT, "fig2_zone_ratios.csv"))
DV = pd.read_csv(os.path.join(OUT, "divergence_pooled.csv"))
rows = []
for _, r in DV.iterrows():
    rows.append(dict(cls=CLS[r.intensity], per=r.window, zone="total", meas=r.measure, ratio=r.ratio, lo=r.lo, hi=r.hi, share=1.0))
for _, r in Z.iterrows():
    rows.append(dict(cls=CLS[r.intensity], per=r.phase, zone=r.tier, meas=r.measure, ratio=r.ratio, lo=r.lo, hi=r.hi, share=r.co2_share))
t3 = pd.DataFrame(rows)
t3["o1"] = t3.cls.map({c: i for i, c in enumerate(CORD)}); t3["o2"] = t3.per.map({p: i for i, p in enumerate(PER)})
t3["o3"] = t3.zone.map({z: i for i, z in enumerate(ZONE)}); t3["o4"] = t3.meas.map({m: i for i, m in enumerate(MEAS)})
t3 = t3.sort_values(["o1", "o2", "o3", "o4"])
T3 = pd.DataFrame({"Cyclone class": t3.cls, "Period": t3.per.map(PER), "Zone": t3.zone.map(ZONE), "Measure": t3.meas.map(MEAS),
                   "Change from normal": t3.ratio - 1, "95% interval, low": t3.lo - 1, "95% interval, high": t3.hi - 1,
                   "Zone share of normal study-region CO2": t3.share})

# ---------------- Table 4
J = pd.read_csv(os.path.join(OUT, "jit_scenarios.csv")); J["W_h"] = J.W_h.astype(str)
J["o1"] = J.window.map({"queue": 0, "event30": 1}); J["o2"] = J.intensity.map(CLS).map({c: i for i, c in enumerate(CORD)})
J["o3"] = J.W_h.map({"12": 0, "24": 1, "48": 2, "upper": 3})
J = J.sort_values(["o1", "o2", "T_h", "o3", "fuel_price"])
T4 = pd.DataFrame({
    "Waiting counted": J.window.map({"queue": "10 days after reopening", "event30": "Whole event (−2 to +30 days)"}),
    "Cyclone class": J.intensity.map(CLS), "Events": J.n_pairs, "Approach length (h)": J.T_h,
    "Wait removed per queued ship (h)": J.W_h.map(lambda w: "all excess waiting" if w == "upper" else int(w)),
    "Fuel price ($/t)": J.fuel_price, "Net excess waiting (ship-hours)": J.net_excess_wait_h.round(0),
    "Excess waiting, 95% low": J.wait_h_lo.round(0), "Excess waiting, 95% high": J.wait_h_hi.round(0),
    "Main-engine fuel saved (t)": J.fuel_saved_t.round(0), "CO2 avoided (t)": J.co2_saved_t.round(0),
    "CO2 avoided, 95% low (t)": J.co2_lo.round(0), "CO2 avoided, 95% high (t)": J.co2_hi.round(0),
    "Generator and boiler fuel moved offshore (t)": J.relocated_hotel_fuel_t.round(0),
    "Fuel cost saved ($)": J.fuel_cost_saved_usd.round(0), "CO2 avoided valued at $51/t ($)": J.scc51_usd.round(0),
    "CO2 avoided valued at $190/t ($)": J.scc190_usd.round(0)})

for n, T in [(2, T2), (3, T3), (4, T4)]:
    T.to_csv(os.path.join(TAB, f"supp_table{n}.csv"), index=False)

# ---------------- workbook
TITLES = {
    2: ("Supplementary Table 2 | Port–cyclone events.",
        "All 344 events at 47 US Gulf and Atlantic ports from 48 tropical cyclones, 2015–2023. Class uses wind at closest approach. "
        "The storm window is the whole days the cyclone was within 500 km of the port; baseline years are the cyclone-free years used for the normal (Methods 4.3–4.4). Blank surge or rainfall cells had no observation."),
    3: ("Supplementary Table 3 | Pooled changes by zone, measure, period and cyclone class.",
        "Change from the no-cyclone normal, pooled over events (observed ÷ expected − 1), with 95% intervals from resampling cyclones. "
        "Study-region rows are the values in Fig. 1; zone rows are the values in Fig. 2 and Extended Data Fig. 5, each zone against its own all-ship normal. "
        "Zone-level values are not available for the whole event."),
    4: ("Supplementary Table 4 | Recovery arrival-management savings under alternative assumptions.",
        "A queued ship that would sail an approach of T hours and then wait W hours instead sails the same distance in T + W hours (Methods 4.6). "
        "\"All excess waiting\" removes each ship's full excess wait. Intervals come from resampling cyclones. "
        "The central case in the text is 10 days after reopening, hurricanes, T = 72 h, W = 24 h, $650 per tonne."),
}
FMT = {2: {}, 3: {"Change from normal": "+0.0%;-0.0%;0.0%", "95% interval, low": "+0.0%;-0.0%;0.0%", "95% interval, high": "+0.0%;-0.0%;0.0%",
                  "Zone share of normal study-region CO2": "0.0%"}, 4: {}}
wb = Workbook(); wb.remove(wb.active)
thin = Side(style="thin", color="8A8984")
for n, T in [(2, T2), (3, T3), (4, T4)]:
    ws = wb.create_sheet(f"Supp Table {n}")
    ws["A1"] = TITLES[n][0]; ws["A1"].font = Font(name="Arial", bold=True, size=11)
    ws["A2"] = TITLES[n][1]; ws["A2"].font = Font(name="Arial", size=9); ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=min(len(T.columns), 12)); ws.row_dimensions[2].height = 42
    for j, c in enumerate(T.columns, 1):
        cell = ws.cell(row=4, column=j, value=c)
        cell.font = Font(name="Arial", bold=True, size=9); cell.alignment = Alignment(wrap_text=True, vertical="bottom")
        cell.fill = PatternFill("solid", fgColor="EFEEEA"); cell.border = Border(bottom=thin)
        width = max(10, min(34, max(len(str(c)) * .55, T[c].astype(str).str.len().max() * 1.1)))
        ws.column_dimensions[get_column_letter(j)].width = width
    ws.row_dimensions[4].height = 40
    for i, rec in enumerate(T.itertuples(index=False), 5):
        for j, v in enumerate(rec, 1):
            cell = ws.cell(row=i, column=j, value=(None if pd.isna(v) else (v.item() if hasattr(v, "item") else v)))
            cell.font = Font(name="Arial", size=9)
            col = T.columns[j - 1]
            if col in FMT[n]:
                cell.number_format = FMT[n][col]
            elif isinstance(v, float) and abs(v) >= 1000:
                cell.number_format = "#,##0"
            elif isinstance(v, (int,)) or (hasattr(v, "dtype") and "int" in str(v.dtype)):
                cell.number_format = "#,##0" if abs(v) >= 1000 else "0"
    ws.freeze_panes = "A5"
wb.save(os.path.join(TAB, "supplementary_tables_2-4.xlsx"))
print(T2.shape, T3.shape, T4.shape)
